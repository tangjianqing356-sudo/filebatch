"""Excel / CSV 的读写与格式识别，所有 Excel 功能共用这一层。

放在这里的都是"怎么把文件读成一张表、怎么把一张表写回去"这类事，
具体某个功能要怎么处理数据，各自在自己的 job 模块里写。
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

CSV_SUFFIXES = {".csv", ".tsv"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
SUPPORTED_SUFFIXES = CSV_SUFFIXES | EXCEL_SUFFIXES

CANDIDATE_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "big5")

# openpyxl 写 xlsx 的硬上限是 1048576 行，留点余量
MAX_ROWS = 1_000_000

# 文件头魔数：用来在真正去读之前就认出"这不是普通 xlsx"
_ZIP_MAGIC = b"PK\x03\x04"          # 正常的 xlsx / xlsm 本质是 zip
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0"   # 老式 .xls，以及**被加密的 xlsx**


@dataclass
class TableData:
    """一张读进来的表。"""

    header: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    source: Path | None = None
    sheet: str = ""

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def col_count(self) -> int:
        return max([len(self.header)] + [len(r) for r in self.rows] or [0]) if (self.header or self.rows) else 0

    def column_index(self, name: str) -> int | None:
        try:
            return self.header.index(name)
        except ValueError:
            return None

    def value(self, row: list[str], column: str) -> str:
        i = self.column_index(column)
        if i is None or i >= len(row):
            return ""
        return row[i]


def sniff_format(path: Path) -> str:
    """看文件头判断真实格式。扩展名会骗人，魔数不会。

    返回 zip / ole2 / text / empty / unknown。
    """
    try:
        with open(path, "rb") as f:
            head = f.read(8)
    except OSError:
        return "unknown"
    if not head:
        return "empty"
    if head.startswith(_ZIP_MAGIC):
        return "zip"
    if head.startswith(_OLE2_MAGIC):
        return "ole2"
    return "text"


def describe_unsupported(path: Path) -> str | None:
    """这个文件能不能处理。返回 None 表示可以，否则返回中文原因。

    刻意不做"静默兼容"：认不出来的格式要明确告诉用户，
    而不是读出一堆乱码让人以为处理成功了。
    """
    suffix = path.suffix.lower()

    if suffix == ".xls":
        return "这是旧版 .xls 格式，暂不支持。请用 Excel 打开后「另存为」.xlsx 再试。"
    if suffix == ".xlsb":
        return "这是二进制 .xlsb 格式，暂不支持。请另存为 .xlsx 再试。"
    if suffix not in SUPPORTED_SUFFIXES:
        return f"不支持的格式：{suffix or '（无扩展名）'}"

    fmt = sniff_format(path)
    if fmt == "empty":
        return "文件是空的"

    if suffix in EXCEL_SUFFIXES:
        if fmt == "ole2":
            # xlsx 被加密后外层是 OLE2 容器，和 .xls 长得一样
            return (
                "这个文件打不开：可能是设置了打开密码的加密文件，"
                "也可能是把 .xls 直接改名成了 .xlsx。请先去掉密码或另存为 .xlsx。"
            )
        if fmt != "zip":
            return "文件内容不像是 Excel 文件，可能已损坏"

    return None


def list_sheets(path: Path) -> list[str]:
    """列出工作表名。CSV 只有一张"表"。"""
    if path.suffix.lower() in CSV_SUFFIXES:
        return ["（CSV 无工作表）"]
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _read_csv_rows(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    text = None
    for enc in CANDIDATE_ENCODINGS:
        try:
            text = raw.decode(enc)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if text is None:
        raise ValueError("CSV 编码无法识别（试过 UTF-8 / GBK / Big5）")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(text.splitlines(), dialect)]


def _read_excel_rows(path: Path, sheet: str | None = None) -> tuple[list[list[str]], str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        if ws is None:
            return [], ""
        rows = [
            ["" if v is None else str(v) for v in row]
            for row in ws.iter_rows(values_only=True)
        ]
        return rows, ws.title
    finally:
        wb.close()


def read_rows(path: Path, sheet: str | None = None) -> tuple[list[list[str]], str]:
    """读成原始的二维列表，不区分表头。返回 (行, 工作表名)。"""
    问题 = describe_unsupported(path)
    if 问题:
        raise ValueError(问题)

    if path.suffix.lower() in CSV_SUFFIXES:
        return _read_csv_rows(path), ""
    return _read_excel_rows(path, sheet)


def read_table(path: Path, has_header: bool = True, sheet: str | None = None) -> TableData:
    """读成 TableData。has_header=False 时表头为空、所有行都算数据。"""
    rows, sheet_name = read_rows(path, sheet)
    if not rows:
        return TableData(source=path, sheet=sheet_name)

    if has_header:
        header = [str(c).strip() for c in rows[0]]
        data = rows[1:]
    else:
        header = []
        data = rows

    return TableData(header=header, rows=data, source=path, sheet=sheet_name)


# Excel 工作表名不允许出现这些字符，也不能超过 31 个字符
_SHEET_ILLEGAL = set('\\/?*[]:')
SHEET_NAME_MAX = 31


def safe_sheet_name(name: str, fallback: str = "结果") -> str:
    """把任意字符串变成合法的 Excel 工作表名。

    工作表名的限制比文件名还多：不能有 \\ / ? * [ ] :，不能超过 31 字符，不能为空。
    拆表功能会拿"地区"这类列值当工作表名，"华东/华南" 直接传给 openpyxl 会报错，
    所以统一在写入层兜住——这是 Excel 格式的约束，不该让每个功能各自去记。
    """
    清洗 = "".join("_" if ch in _SHEET_ILLEGAL else ch for ch in str(name))
    清洗 = 清洗.strip().strip("'")          # 首尾空格和单引号 Excel 也不接受
    清洗 = 清洗[:SHEET_NAME_MAX]
    return 清洗 or fallback


def is_blank_row(row: list[str]) -> bool:
    return not any(str(c).strip() for c in row)


def write_table(
    data: TableData,
    target: Path,
    output_format: str = "xlsx",
    sheet_name: str = "结果",
) -> None:
    """把一张表写出去。调用方负责保证 target 不会覆盖已有文件。"""
    target.parent.mkdir(parents=True, exist_ok=True)

    if output_format == "csv":
        with open(target, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            if data.header:
                writer.writerow(data.header)
            writer.writerows(data.rows)
        return

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = safe_sheet_name(sheet_name)
    if data.header:
        ws.append(data.header)
    for row in data.rows:
        ws.append(row)
    wb.save(target)


def write_sheets(
    target: Path,
    sheets: list[tuple[str, TableData]],
) -> None:
    """一次写出多张工作表（两表对比要用）。"""
    from openpyxl import Workbook

    target.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    用过: set[str] = set()
    for name, data in sheets:
        标题 = safe_sheet_name(name, "表")
        # 同名工作表 Excel 也不接受，清洗后可能撞名，这里再避让一次
        if 标题 in 用过:
            for n in range(1, 1000):
                候选 = safe_sheet_name(f"{标题[:SHEET_NAME_MAX - 3]}_{n}", "表")
                if 候选 not in 用过:
                    标题 = 候选
                    break
        用过.add(标题)
        ws = wb.create_sheet(title=标题)
        if data.header:
            ws.append(data.header)
        for row in data.rows:
            ws.append(row)
    if not wb.sheetnames:
        wb.create_sheet(title="结果")
    wb.save(target)
