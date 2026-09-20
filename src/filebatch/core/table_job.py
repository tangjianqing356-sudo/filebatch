"""Excel / CSV 批量合并成一个文件。

故意不用 pandas：这里只是把多个表的行拼起来，用 openpyxl + 标准库 csv 就够了，
能省掉 numpy 那 60MB 和 PyInstaller 打包时的一堆坑。

列对齐策略：有表头时按**表头名**对齐（缺的列补空），这是最不容易丢数据的做法；
无表头时按位置对齐，短的行补空。
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from .result import JobReport, PlannedAction
from .safety import unique_path

CSV_SUFFIXES = {".csv"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
SUPPORTED_SUFFIXES = CSV_SUFFIXES | EXCEL_SUFFIXES

CANDIDATE_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "big5")

MAX_ROWS = 1_000_000   # openpyxl 的 xlsx 上限是 1048576 行


@dataclass
class MergeOptions:
    has_header: bool = True
    add_source_column: bool = True          # 加一列记录数据来自哪个文件，排查问题很有用
    source_column_name: str = "来源文件"
    output_format: str = "xlsx"             # xlsx 或 csv
    sheet_name: str = "合并结果"

    def validate(self) -> str | None:
        if self.output_format not in ("xlsx", "csv"):
            return "输出格式只支持 xlsx 或 csv"
        if self.add_source_column and not self.source_column_name.strip():
            return "来源列的列名不能为空"
        return None


def _read_csv(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    text = None
    for enc in CANDIDATE_ENCODINGS:
        try:
            text = raw.decode(enc)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if text is None:
        raise ValueError("CSV 编码无法识别（试过 UTF-8 / GBK / Big5）")

    # 自动识别分隔符，逗号和制表符都常见
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(text.splitlines(), dialect)]


def _read_excel(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        if ws is None:
            return []
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
        return rows
    finally:
        wb.close()


def read_table(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix in CSV_SUFFIXES:
        return _read_csv(path)
    if suffix in EXCEL_SUFFIXES:
        return _read_excel(path)
    raise ValueError(f"不支持的表格格式：{suffix}")


def align_rows(
    tables: list[tuple[str, list[list[str]]]],
    options: MergeOptions,
) -> tuple[list[str], list[list[str]]]:
    """把多个表对齐成统一的表头和行。纯函数，方便测试。

    :param tables: [(来源文件名, 行数据), ...]
    :return: (表头, 所有数据行)
    """
    if options.has_header:
        header: list[str] = []
        seen: set[str] = set()
        for _, rows in tables:
            if not rows:
                continue
            for col in rows[0]:
                name = str(col).strip()
                if name and name not in seen:
                    seen.add(name)
                    header.append(name)

        merged: list[list[str]] = []
        for source, rows in tables:
            if len(rows) < 2:
                continue
            own_header = [str(c).strip() for c in rows[0]]
            index_of = {name: i for i, name in enumerate(own_header)}
            for row in rows[1:]:
                if not any(str(c).strip() for c in row):
                    continue   # 跳过整行空白
                out = []
                for name in header:
                    i = index_of.get(name)
                    out.append(str(row[i]) if i is not None and i < len(row) else "")
                if options.add_source_column:
                    out.append(source)
                merged.append(out)

        if options.add_source_column:
            header = header + [options.source_column_name]
        return header, merged

    # 无表头：按位置对齐
    width = max((len(r) for _, rows in tables for r in rows), default=0)
    merged = []
    for source, rows in tables:
        for row in rows:
            if not any(str(c).strip() for c in row):
                continue
            out = [str(row[i]) if i < len(row) else "" for i in range(width)]
            if options.add_source_column:
                out.append(source)
            merged.append(out)
    header = []
    return header, merged


def plan_merge(files: list[Path], options: MergeOptions, output_file: Path) -> list[PlannedAction]:
    """计划阶段会真的读一遍表，预览里能看到每个文件有多少行、多少列。"""
    actions: list[PlannedAction] = []

    for src in files:
        if src.suffix.lower() not in SUPPORTED_SUFFIXES:
            hint = "（.xls 旧格式请先用 Excel 另存为 .xlsx）" if src.suffix.lower() == ".xls" else ""
            actions.append(
                PlannedAction(src, None, "合并表格",
                              skip_reason=f"不支持的格式：{src.suffix}{hint}")
            )
            continue
        try:
            rows = read_table(src)
            if not rows:
                actions.append(PlannedAction(src, None, "合并表格", skip_reason="表格是空的"))
                continue
            data_rows = max(0, len(rows) - 1) if options.has_header else len(rows)
            actions.append(
                PlannedAction(src, output_file, "合并表格",
                              f"{data_rows} 行数据，{len(rows[0])} 列")
            )
        except PermissionError:
            actions.append(PlannedAction(src, None, "合并表格", skip_reason="没有权限读取该文件"))
        except Exception as e:
            actions.append(PlannedAction(src, None, "合并表格", skip_reason=f"读取失败：{e}"))

    return actions


def execute(
    actions: list[PlannedAction],
    options: MergeOptions,
    output_file: Path,
    on_progress: Callable[[int, int, str], None] | None = None,
    should_cancel: Callable[[], bool] | None = None,
) -> JobReport:
    """把多个表合成一个。

    中断后**已经读进来的表仍会写出**，用户拿到的是一份完整可用的结果文件。

    :param on_progress: 每处理完一项回调 (已完成数, 总数, 当前文件名)，供界面显示进度
    :param should_cancel: 返回 True 表示用户要求中断。
        中断只影响**尚未执行**的项目，已经处理完的原样保留。
    """
    report = JobReport("Excel / CSV 批量合并")
    tables: list[tuple[str, list[list[str]]]] = []
    total = len(actions)

    for done, act in enumerate(actions):
        if should_cancel is not None and should_cancel():
            for rest in actions[done:]:
                report.add_skip(rest.source, "用户中断任务，此项未合并")
            break
        if not act.will_run:
            report.add_skip(act.source, act.skip_reason or "未说明")
            continue
        try:
            rows = read_table(act.source)
            tables.append((act.source.name, rows))
            report.add_ok(act.source, act.target, act.note)
        except Exception as e:
            report.add_fail(act.source, f"读取失败：{type(e).__name__} {e}")

        if on_progress is not None:
            on_progress(done + 1, total, act.source.name)

    if not tables:
        return report.finish()

    try:
        header, merged = align_rows(tables, options)
        if len(merged) > MAX_ROWS:
            for item in report.items:
                item.ok = False
                item.message = f"合并后共 {len(merged)} 行，超过 {MAX_ROWS} 行上限"
            return report.finish()

        final = unique_path(output_file)
        final.parent.mkdir(parents=True, exist_ok=True)

        if options.output_format == "csv":
            with open(final, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                if header:
                    writer.writerow(header)
                writer.writerows(merged)
        else:
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = options.sheet_name
            if header:
                ws.append(header)
            for row in merged:
                ws.append(row)
            wb.save(final)

        for item in report.items:
            if item.ok:
                item.target = final
        report.job_name += f"：共 {len(merged)} 行 -> {final.name}"
    except Exception as e:
        for item in report.items:
            if item.ok:
                item.ok = False
                item.message = f"写入合并结果失败：{type(e).__name__} {e}"

    return report.finish()
