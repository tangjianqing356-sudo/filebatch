"""打包产物的完整验收：每个功能各走一遍真实流程。

走的是**真实界面对象**，不是直接调 core：
添加文件 → 设参数 → 生成预览 → 开始执行 → 等后台线程 → 核对磁盘产物 → 日志。

之所以要在打包后再跑一遍，是因为"开发环境能跑"和"打出来能跑"是两回事：
少一个 Pillow 插件、少一个 Qt plugin，只有在冻结后的环境里才暴露。

用法：FileBatchTool --acceptance [--open-output]
"""
from __future__ import annotations

import csv
import sys
import tempfile
import time
import traceback
from pathlib import Path

结果: list[tuple[str, bool, str]] = []


def _静音弹窗() -> None:
    from filebatch.ui.pages import base_page as bp

    bp.show_info = lambda *a, **k: None
    bp.show_error = lambda *a, **k: None
    bp.confirm_dangerous = lambda *a, **k: True


def _等执行完(app, page, 超时秒: float = 60) -> None:
    from PySide6.QtCore import QDeadlineTimer

    worker = page._worker
    t0 = time.time()
    while worker is not None and worker.isRunning() and time.time() - t0 < 超时秒:
        app.processEvents()
        worker.wait(QDeadlineTimer(20))
    for _ in range(10):
        app.processEvents()


def _主功能(序号: int):
    """左边导航第 序号 项那一页。"""
    def 定位(win):
        page = win.page_at(序号)
        win.nav.setCurrentRow(序号)
        return page
    return 定位


def _工具箱(标签序号: int):
    """Excel 工具箱里的第 标签序号 个标签页。"""
    def 定位(win):
        from filebatch.ui.main_window import EXCEL_TOOLBOX_INDEX

        toolbox = win.page_at(EXCEL_TOOLBOX_INDEX)
        win.nav.setCurrentRow(EXCEL_TOOLBOX_INDEX)
        toolbox.tabs.setCurrentIndex(标签序号)
        return toolbox.sub_page(标签序号)
    return 定位


def _等列名(app, page, 超时秒: float = 30) -> None:
    """选列的页面要等后台把表头读回来。不选列的页面直接返回。"""
    worker = getattr(page, "_inspect_worker", None)
    if worker is None:
        return
    t0 = time.time()
    while worker.isRunning() and time.time() - t0 < 超时秒:
        app.processEvents()
        worker.wait(20)
    for _ in range(10):
        app.processEvents()


def _跑一个功能(app, win, 定位, 名称: str, 准备, 设参数, 校验, open_output: bool) -> None:
    try:
        page = 定位(win)
        for _ in range(3):
            app.processEvents()

        tmp = Path(tempfile.mkdtemp(prefix="fb_acc_"))
        源, 期望数 = 准备(tmp)

        page.file_list.add_paths_sync([源])
        assert page.file_list.files(), "文件清单是空的"
        assert not page.btn_execute.isEnabled(), "预览之前不该能执行"
        # 选列的页面是后台线程读的表头，得等它回来，设参数时才勾得上列
        _等列名(app, page)

        设参数(page, tmp)

        assert page.generate_preview() is True, "生成预览失败"
        assert page.btn_execute.isEnabled(), "预览通过后执行按钮应该可用"

        assert page.start_execute() is True, "启动执行失败"
        _等执行完(app, page)

        report = page._last_report
        assert report is not None, "没拿到执行结果"
        assert report.failed == 0, f"有 {report.failed} 项失败：{report.summary()}"
        assert report.succeeded == 期望数, f"期望成功 {期望数}，实际 {report.succeeded}"

        校验(tmp, page)

        assert page.log.toPlainText().strip(), "日志是空的"
        assert page.btn_save_log.isEnabled(), "保存日志按钮应该可用"
        assert page.btn_open_output.isEnabled(), "打开输出目录按钮应该可用"
        assert "本次执行结果" in page.box_preview.title(), "执行后应切到结果态"

        打开 = ""
        if open_output:
            from PySide6.QtCore import QUrl
            from PySide6.QtGui import QDesktopServices

            目标 = page._output_path
            if 目标 and 目标.is_file():
                目标 = 目标.parent
            ok = QDesktopServices.openUrl(QUrl.fromLocalFile(str(目标)))
            打开 = f"，已打开输出目录={ok}"

        结果.append((名称, True, f"{report.succeeded} 项成功，日志 {len(page.log.toPlainText())} 字{打开}"))
    except Exception as e:
        结果.append((名称, False, f"{type(e).__name__}: {e}"))


# ---------------- 各功能的样例数据与校验 ----------------

def _备重命名(tmp: Path):
    d = tmp / "照片"
    d.mkdir()
    for i in range(6):
        (d / f"IMG_{i:04d}.jpg").write_bytes(b"x" * 1024)
    return d, 6


def _设重命名(page, tmp):
    page._output_path = tmp / "输出"
    page.ed_base.setText("旅行照片")
    page.cb_number.setCurrentIndex(1)
    page.sp_digits.setValue(3)


def _校验重命名(tmp, page):
    out = tmp / "输出"
    names = sorted(p.name for p in out.iterdir())
    assert names[0] == "旅行照片_001.jpg", names[:3]
    assert len(names) == 6
    assert (tmp / "照片" / "IMG_0000.jpg").exists(), "原文件必须保留"


def _备分类(tmp: Path):
    d = tmp / "杂乱"
    d.mkdir()
    for n in ["图.jpg", "表.xlsx", "文.txt", "包.zip", "片.mp4"]:
        (d / n).write_bytes(b"x" * 512)
    return d, 5


def _设分类(page, tmp):
    page._output_path = tmp / "整理后"


def _校验分类(tmp, page):
    out = tmp / "整理后"
    for folder, name in [("图片", "图.jpg"), ("表格", "表.xlsx"), ("文档", "文.txt"),
                         ("压缩包", "包.zip"), ("视频", "片.mp4")]:
        assert (out / folder / name).exists(), f"{folder}/{name} 不存在"
    assert (tmp / "杂乱" / "图.jpg").exists(), "默认复制，原文件必须保留"


def _备图片(tmp: Path):
    from PIL import Image

    d = tmp / "原图"
    d.mkdir()
    for i in range(4):
        Image.new("RGB", (1600, 1200), (90, 140, 200)).save(d / f"DSC_{i}.jpg")
    return d, 4


def _设图片(page, tmp):
    from filebatch.core.image_job import ResizeMode

    page._output_path = tmp / "输出"
    page.ed_base.setText("产品图")
    page.cb_number.setCurrentIndex(1)
    page.cb_resize.setCurrentIndex(page.cb_resize.findData(ResizeMode.MAX_SIDE))
    page.sp_value.setValue(800)


def _校验图片(tmp, page):
    from PIL import Image

    out = tmp / "输出"
    assert len(list(out.iterdir())) == 4
    with Image.open(out / "产品图_001.jpg") as im:
        assert im.size == (800, 600), f"缩放结果不对：{im.size}"


def _备表格(tmp: Path):
    d = tmp / "报表"
    d.mkdir()
    for i, m in enumerate(["一月", "二月", "三月"], start=1):
        with open(d / f"{m}.csv", "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(["姓名", "金额"] if i % 2 else ["金额", "姓名"])
            w.writerow([f"张{i}", str(i * 100)] if i % 2 else [str(i * 100), f"张{i}"])
    return d, 3


def _设表格(page, tmp):
    page._output_path = tmp / "合并结果.xlsx"


def _校验表格(tmp, page):
    from openpyxl import load_workbook

    ws = load_workbook(tmp / "合并结果.xlsx").active
    assert [c.value for c in ws[1]] == ["姓名", "金额", "来源文件"], [c.value for c in ws[1]]
    assert ws.max_row == 4, f"应有 1 行表头 + 3 行数据，实际 {ws.max_row}"


def _备PDF(tmp: Path):
    from pypdf import PdfWriter

    d = tmp / "合同"
    d.mkdir()
    w = PdfWriter()
    for _ in range(4):
        w.add_blank_page(width=595, height=842)
    with open(d / "采购合同.pdf", "wb") as f:
        w.write(f)
    return d, 4


def _设PDF(page, tmp):
    page._output_path = tmp / "输出"


def _校验PDF(tmp, page):
    from pypdf import PdfReader

    out = tmp / "输出"
    names = sorted(p.name for p in out.iterdir())
    assert names == [f"采购合同_第{i}页.pdf" for i in range(1, 5)], names
    assert len(PdfReader(str(out / "采购合同_第1页.pdf")).pages) == 1


def _备文本(tmp: Path):
    d = tmp / "文案"
    d.mkdir()
    for i in range(3):
        (d / f"doc{i}.txt").write_text("旧公司名 出品，旧公司名 版权所有", encoding="utf-8")
    return d, 3


def _设文本(page, tmp):
    page._output_path = tmp / "输出"
    page.ed_find.setText("旧公司名")
    page.ed_replace.setText("新公司名")


def _校验文本(tmp, page):
    out = tmp / "输出"
    内容 = (out / "doc0.txt").read_text(encoding="utf-8")
    assert 内容 == "新公司名 出品，新公司名 版权所有", 内容
    原 = (tmp / "文案" / "doc0.txt").read_text(encoding="utf-8")
    assert "旧公司名" in 原, "原文件必须不变"


# ---------------- Excel 工具箱 ----------------

def _写xlsx(path: Path, rows: list[list]) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def _备去重(tmp: Path):
    """20 个正常文件 + 1 个缺列的。

    列下拉只读前 20 个文件，所以「手机」会出现在可选列里；
    第 21 个文件没有这一列，必须在**预览阶段**就被点名跳过，
    而不是等到 execute 才一个个失败。这是本轮专门要验的行为。
    """
    from filebatch.ui.excel_toolbox.columns import MAX_INSPECT_FILES

    d = tmp / "去重源"
    d.mkdir()
    for i in range(MAX_INSPECT_FILES):
        _写xlsx(d / f"{i:02d}_客户.xlsx", [
            ["姓名", "手机", "城市"],
            ["张三", "13800000001", "北京"],
            ["李四", "13800000002", "上海"],
            ["张三", "13800000001", "北京"],      # 完全重复
        ])
    _写xlsx(d / "99_缺列的表.xlsx", [["姓名"], ["赵六"], ["赵六"]])
    return d, MAX_INSPECT_FILES


def _设去重(page, tmp):
    page._output_path = tmp / "输出"
    页面列 = [page.list_keys.item(i).text() for i in range(page.list_keys.count())]
    assert "手机" in 页面列, f"前 20 个文件都有手机列，下拉里应该有：{页面列}"
    page.list_keys.set_checked(["姓名", "手机"])


def _校验去重(tmp, page):
    from openpyxl import load_workbook

    from filebatch.ui.excel_toolbox.columns import MAX_INSPECT_FILES

    ws = load_workbook(tmp / "输出" / "00_客户_去重.xlsx").active
    assert ws.max_row == 3, f"1 行表头 + 2 行去重后数据，实际 {ws.max_row}"
    assert [c.value for c in ws[1]] == ["姓名", "手机", "城市"]
    assert (tmp / "去重源" / "00_客户.xlsx").exists(), "原文件必须保留"

    # 缺列的文件必须在**执行前**就被查出来并写进日志
    报告 = page._header_report
    assert 报告 is not None, "预览阶段没有跑全量表头检查"
    assert 报告.total == MAX_INSPECT_FILES + 1, f"要检查全部文件，实际只检查了 {报告.total} 个"
    assert [p.name for p in 报告.blocked] == ["99_缺列的表.xlsx"], [p.name for p in 报告.blocked]
    assert "99_缺列的表.xlsx" in page.log.toPlainText(), "日志里没点名缺列的文件"
    assert not (tmp / "输出" / "99_缺列的表_去重.xlsx").exists(), "缺列的文件不该产出结果"


def _备对比(tmp: Path):
    d = tmp / "对比源"
    d.mkdir()
    # 文件名决定清单顺序：A 表在前，B 表在后
    _写xlsx(d / "1_上月.xlsx", [
        ["工号", "姓名", "金额"],
        ["001", "张三", "100"],
        ["002", "李四", "200"],
        ["003", "王五", "300"],
    ])
    _写xlsx(d / "2_本月.xlsx", [
        ["工号", "姓名", "金额"],
        ["001", "张三", "100"],      # 完全一样
        ["002", "李四", "250"],      # 金额变了
        ["004", "赵六", "400"],      # 新增
    ])
    return d, 1


def _设对比(page, tmp):
    page._output_path = tmp / "对比结果.xlsx"
    # 明确选 A / B，不靠清单顺序
    a = next(f for f in page.file_list.files() if f.name.startswith("1_"))
    b = next(f for f in page.file_list.files() if f.name.startswith("2_"))
    page.cb_a.setCurrentIndex(page.cb_a.findData(str(a)))
    page.cb_b.setCurrentIndex(page.cb_b.findData(str(b)))
    page.list_keys.set_checked(["工号"])


def _校验对比(tmp, page):
    from openpyxl import load_workbook

    a, b = page.selected_pair()
    assert a.name.startswith("1_") and b.name.startswith("2_"), (a.name, b.name)

    wb = load_workbook(tmp / "对比结果.xlsx")
    for name in ("仅A表有", "仅B表有", "内容不同", "汇总"):
        assert name in wb.sheetnames, f"缺少工作表 {name}：{wb.sheetnames}"
    assert wb["仅A表有"].max_row == 2, "王五 只在 A 表"
    assert wb["仅B表有"].max_row == 2, "赵六 只在 B 表"
    assert wb["内容不同"].max_row == 2, "李四 的金额不同"


def _备清洗(tmp: Path):
    d = tmp / "清洗源"
    d.mkdir()
    _写xlsx(d / "名单.xlsx", [
        ["姓名", "备注", "编号"],
        ["  张三  ", "说明", "001"],
        ["", "", ""],                       # 空行
        ["李四", "多个   空格", "0012"],
    ])
    return d, 1


def _设清洗(page, tmp):
    page._output_path = tmp / "输出"


def _校验清洗(tmp, page):
    from openpyxl import load_workbook

    ws = load_workbook(tmp / "输出" / "名单_清洗.xlsx").active
    assert ws.max_row == 3, f"空行应被删掉，实际 {ws.max_row} 行"
    assert ws["A2"].value == "张三", f"首尾空格应被去掉：{ws['A2'].value!r}"
    assert ws["B3"].value == "多个 空格", f"连续空格应并成一个：{ws['B3'].value!r}"
    # 最容易踩的坑：以 0 开头的编号不能被当成数字
    assert ws["C2"].value == "001", f"编号被改掉了：{ws['C2'].value!r}"


def _备拆表(tmp: Path):
    d = tmp / "拆表源"
    d.mkdir()
    _写xlsx(d / "花名册.xlsx", [
        ["姓名", "部门"],
        ["张三", "研发"],
        ["李四", "销售"],
        ["王五", "研发"],
        ["赵六", "行政/后勤"],       # 文件名和工作表名里都不能直接用的字符
    ])
    return d, 3


def _设拆表(page, tmp):
    page._output_path = tmp / "输出"
    page.cb_column.setCurrentText("部门")


def _校验拆表(tmp, page):
    from openpyxl import load_workbook

    out = tmp / "输出"
    names = sorted(p.name for p in out.iterdir())
    assert len(names) == 3, names
    assert "花名册_研发.xlsx" in names, names
    assert not any("/" in n for n in names), f"文件名里不该有斜杠：{names}"
    ws = load_workbook(out / "花名册_研发.xlsx").active
    assert ws.max_row == 3, "1 行表头 + 2 个研发的人"


def _备格式(tmp: Path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font

    d = tmp / "格式源"
    d.mkdir()

    # 一张字体混排的表：默认参数跑完，这些字体必须原样还在。
    # 「格式统一」不该顺手把人家表里的字体全换成我们的默认值。
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r in [["姓名", "金额"], ["张三", 100], ["李四", 200]]:
        ws.append(r)
    ws["A1"].font = Font(name="Times New Roman", size=13)
    ws["A2"].font = Font(name="宋体", size=14)
    ws["A3"].font = Font(name="Arial", size=9)
    # 对齐也是左/中/右混排：默认跑完必须原样
    ws["A1"].alignment = Alignment(horizontal="right")
    ws["A2"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws["A3"].alignment = Alignment(horizontal="center")
    wb.save(d / "报表.xlsx")

    _写xlsx(d / "报表2.xlsx", [["项目", "数量"], ["A", 1]])

    # 一张"什么都有"的表：图表 + 图片 + 条件格式 + 表格对象。
    # 交付说明里写的是"验收样例已验证可保留"，这就是那个验收样例——
    # 必须在**冻结产物**里核对，不能只在 pytest 里过。
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from PIL import Image as PILImage

    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    for r in [["月份", "销量"], ["一月", 10], ["二月", 20], ["三月", 30]]:
        ws.append(r)

    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    ws.add_chart(chart, "E2")

    png = d / "_临时图.png"
    PILImage.new("RGB", (40, 40), (90, 140, 200)).save(png)
    ws.add_image(XLImage(str(png)), "H2")

    ws.conditional_formatting.add("B2:B10", CellIsRule(
        operator="greaterThan", formula=["15"],
        fill=PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")))

    表 = Table(displayName="销量表", ref="A1:B4")
    表.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(表)

    wb.save(d / "全都有.xlsx")
    png.unlink()
    return d, 3


def _设格式(page, tmp):
    page._output_path = tmp / "输出"


def _校验格式(tmp, page):
    from openpyxl import load_workbook

    ws = load_workbook(tmp / "输出" / "报表_格式化.xlsx").active
    assert ws["A1"].font.bold is True, "表头应该加粗"
    assert ws.freeze_panes == "A2", "表头应该冻结"
    assert ws.column_dimensions["A"].width > 0, "列宽应该被设过"
    assert (tmp / "格式源" / "报表.xlsx").exists(), "原文件必须保留"

    # 默认参数 = 保留原字体。冻结环境里也必须是这个行为。
    assert page.selected_font() is None, "默认不该选中任何字体"
    assert page.selected_size() is None, "默认不该改字号"
    字体 = [(ws[c].font.name, ws[c].font.size) for c in ("A1", "A2", "A3")]
    assert 字体 == [("Times New Roman", 13), ("宋体", 14), ("Arial", 9)], 字体

    # 默认参数 = 保留原对齐
    assert page.selected_align(page.cb_body_align) is None, "默认不该改正文对齐"
    assert page.selected_align(page.cb_head_align) is None, "默认不该改表头对齐"
    对齐 = [ws[c].alignment.horizontal for c in ("A1", "A2", "A3")]
    assert 对齐 == ["right", "left", "center"], 对齐
    assert ws["A2"].alignment.vertical == "top", "垂直对齐被误伤"
    assert ws["A2"].alignment.wrap_text is True, "自动换行被误伤"

    # 顺带确认数字没被变成文本
    assert isinstance(ws["B2"].value, int), f"金额被改成了 {type(ws['B2'].value).__name__}"

    # 能力边界的验收样例：图表 / 图片 / 条件格式 / 表格对象
    全 = load_workbook(tmp / "输出" / "全都有_格式化.xlsx")
    w = 全["数据"]
    缺 = []
    if not getattr(w, "_charts", None):
        缺.append("图表")
    if not getattr(w, "_images", None):
        缺.append("图片")
    if not list(w.conditional_formatting):
        缺.append("条件格式")
    if not getattr(w, "_tables", None):
        缺.append("表格对象")
    assert not 缺, f"验收样例里这些没保住：{'、'.join(缺)}"
    assert w["B2"].value == 10 and isinstance(w["B2"].value, int)


# (定位函数, 名称, 准备, 设参数, 校验)
任务 = [
    (_主功能(0), "批量重命名 / 编号", _备重命名, _设重命名, _校验重命名),
    (_主功能(1), "文件分类整理", _备分类, _设分类, _校验分类),
    (_主功能(2), "图片批处理", _备图片, _设图片, _校验图片),
    (_工具箱(0), "Excel 工具箱 · 合并", _备表格, _设表格, _校验表格),
    (_工具箱(1), "Excel 工具箱 · 去重", _备去重, _设去重, _校验去重),
    (_工具箱(2), "Excel 工具箱 · 两表对比", _备对比, _设对比, _校验对比),
    (_工具箱(3), "Excel 工具箱 · 数据清洗", _备清洗, _设清洗, _校验清洗),
    (_工具箱(4), "Excel 工具箱 · 批量拆表", _备拆表, _设拆表, _校验拆表),
    (_工具箱(5), "Excel 工具箱 · 格式统一", _备格式, _设格式, _校验格式),
    (_主功能(4), "PDF 拆分", _备PDF, _设PDF, _校验PDF),
    (_主功能(5), "文本查找替换", _备文本, _设文本, _校验文本),
]


def run(open_output: bool = False) -> int:
    from filebatch.ui.app import create_app
    from filebatch.ui.main_window import MainWindow

    _静音弹窗()
    app = create_app([])
    win = MainWindow()
    win.show()
    for _ in range(10):
        app.processEvents()

    print("=" * 60)
    print(f"打包产物完整验收（平台插件 = {app.platformName()}）")
    print("=" * 60)

    for i, (定位, 名称, 准备, 设参数, 校验) in enumerate(任务):
        # 只对第一个功能真的去打开输出目录，避免弹出一堆访达窗口
        _跑一个功能(app, win, 定位, 名称, 准备, 设参数, 校验,
                   open_output and i == 0)

    失败 = 0
    for 名称, ok, detail in 结果:
        print(f"{'[通过]' if ok else '[失败]'} {名称}")
        print(f"       {detail}")
        if not ok:
            失败 += 1
    print("-" * 60)
    print(f"共 {len(结果)} 个功能，失败 {失败} 个")

    win.close()
    return 1 if 失败 else 0


if __name__ == "__main__":
    try:
        sys.exit(run("--open-output" in sys.argv))
    except Exception:
        traceback.print_exc()
        sys.exit(2)
