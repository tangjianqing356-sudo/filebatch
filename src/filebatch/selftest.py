"""打包产物自检。

打包最容易出的问题是"开发环境好好的，打出来就少东西"：
Pillow 的格式插件是运行时按需导入的、Qt 的 platform 插件是动态加载的，
静态分析都看不到。所以让打包后的二进制自己跑一遍，比在开发环境里猜靠谱。

用法：FileBatchTool --selftest
"""
from __future__ import annotations

import sys
import tempfile
import traceback
from pathlib import Path

结果: list[tuple[str, bool, str]] = []


def 检查(名称: str, fn) -> None:
    try:
        detail = fn() or ""
        结果.append((名称, True, str(detail)))
    except Exception as e:
        结果.append((名称, False, f"{type(e).__name__}: {e}"))


def _pillow() -> str:
    from PIL import Image
    import PIL

    tmp = Path(tempfile.mkdtemp())
    formats = []
    for ext, mode in (("jpg", "RGB"), ("png", "RGBA"), ("webp", "RGB"), ("bmp", "RGB"), ("gif", "P")):
        p = tmp / f"t.{ext}"
        Image.new(mode, (20, 20)).save(p)
        with Image.open(p) as im:
            im.load()
        formats.append(ext)
    return f"Pillow {PIL.__version__}，可读写 {'/'.join(formats)}"


def _openpyxl() -> str:
    import openpyxl
    from openpyxl import Workbook, load_workbook

    tmp = Path(tempfile.mkdtemp()) / "t.xlsx"
    wb = Workbook()
    wb.active.append(["中文列", 1])
    wb.save(tmp)
    ws = load_workbook(tmp).active
    assert ws["A1"].value == "中文列"
    return f"openpyxl {openpyxl.__version__}，中文读写正常"


def _pypdf() -> str:
    import pypdf
    from pypdf import PdfReader, PdfWriter

    tmp = Path(tempfile.mkdtemp()) / "t.pdf"
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    w.add_blank_page(width=200, height=200)
    with open(tmp, "wb") as f:
        w.write(f)
    assert len(PdfReader(str(tmp)).pages) == 2
    return f"pypdf {pypdf.__version__}"


def _qt() -> str:
    import PySide6
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance() or QApplication([])
    plat = app.platformName()
    return f"PySide6 {PySide6.__version__}，platform 插件 = {plat}"


def _界面构建() -> str:
    from filebatch.ui.app import create_app
    from filebatch.ui.main_window import MainWindow

    app = create_app([])
    win = MainWindow()
    win.show()
    for _ in range(5):
        app.processEvents()
    页数 = len(win.pages)
    首页 = win.stack.currentWidget() is win.pages[0]
    子页数 = 0
    for i in range(页数):
        page = win.page_at(i)
        # Excel 工具箱里的子页面是懒加载的，冻结环境里也得挨个建一遍才算验过
        建子页 = getattr(page, "sub_page", None)
        if 建子页 is not None:
            for j in range(page.tabs.count()):
                建子页(j)
                子页数 += 1
    win.close()
    return f"{页数} 个功能页 + {子页数} 个 Excel 子页全部可创建，首页正确显示={首页}"


def _中文路径() -> str:
    tmp = Path(tempfile.mkdtemp()) / "中文 目录 with space" / "深一层"
    tmp.mkdir(parents=True)
    f = tmp / "中文文件名 带空格 😀.txt"
    f.write_text("内容", encoding="utf-8")
    assert f.read_text(encoding="utf-8") == "内容"
    return "中文 / 空格 / emoji 路径读写正常"


def _完整流程() -> str:
    from filebatch.core import rename_job
    from filebatch.core.naming import NameRule, NumberPosition

    tmp = Path(tempfile.mkdtemp())
    src = tmp / "源"
    src.mkdir()
    files = []
    for i in range(5):
        p = src / f"中文文件{i}.txt"
        p.write_text("x", encoding="utf-8")
        files.append(p)

    out = tmp / "输出"
    rule = NameRule(base_name="照片", number_position=NumberPosition.SUFFIX, number_digits=3)
    report = rename_job.execute(rename_job.plan_rename(files, rule, out))
    assert report.succeeded == 5, f"只成功了 {report.succeeded}"
    assert (out / "照片_001.txt").exists()
    assert all(f.exists() for f in files), "原文件应保留"
    return "计划→执行→结果 全链路 5/5 成功"


def run() -> int:
    检查("Qt 运行时", _qt)
    检查("Pillow 图片格式", _pillow)
    检查("openpyxl 表格", _openpyxl)
    检查("pypdf", _pypdf)
    检查("中文/空格/emoji 路径", _中文路径)
    检查("核心流程", _完整流程)
    检查("界面构建", _界面构建)

    print("=" * 56)
    print("打包产物自检")
    print("=" * 56)
    失败 = 0
    for 名称, ok, detail in 结果:
        mark = "[通过]" if ok else "[失败]"
        print(f"{mark} {名称}")
        if detail:
            print(f"       {detail}")
        if not ok:
            失败 += 1
    print("-" * 56)
    print(f"共 {len(结果)} 项，失败 {失败} 项")
    return 1 if 失败 else 0


if __name__ == "__main__":
    try:
        sys.exit(run())
    except Exception:
        traceback.print_exc()
        sys.exit(2)
