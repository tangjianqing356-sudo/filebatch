import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

from filebatch.core import table_job
from filebatch.core.table_job import MergeOptions, align_rows, read_table


def write_csv(path: Path, rows, encoding="utf-8") -> Path:
    with open(path, "w", encoding=encoding, newline="") as f:
        csv.writer(f).writerows(rows)
    return path


def write_xlsx(path: Path, rows) -> Path:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


# ---------- 纯函数：对齐 ----------

def test_有表头时按列名对齐():
    tables = [
        ("一月.csv", [["姓名", "金额"], ["张三", "100"]]),
        ("二月.csv", [["姓名", "金额"], ["李四", "200"]]),
    ]
    header, rows = align_rows(tables, MergeOptions(add_source_column=False))
    assert header == ["姓名", "金额"]
    assert rows == [["张三", "100"], ["李四", "200"]]


def test_列顺序不同也能按列名正确对齐():
    tables = [
        ("a.csv", [["姓名", "金额"], ["张三", "100"]]),
        ("b.csv", [["金额", "姓名"], ["200", "李四"]]),   # 顺序颠倒
    ]
    header, rows = align_rows(tables, MergeOptions(add_source_column=False))
    assert header == ["姓名", "金额"]
    assert rows == [["张三", "100"], ["李四", "200"]], "必须按列名对齐，不能按位置"


def test_某个文件缺列时补空而不是丢数据():
    tables = [
        ("a.csv", [["姓名", "金额", "备注"], ["张三", "100", "VIP"]]),
        ("b.csv", [["姓名", "金额"], ["李四", "200"]]),
    ]
    header, rows = align_rows(tables, MergeOptions(add_source_column=False))
    assert header == ["姓名", "金额", "备注"]
    assert rows == [["张三", "100", "VIP"], ["李四", "200", ""]]


def test_来源列默认会加上():
    tables = [("一月.csv", [["姓名"], ["张三"]])]
    header, rows = align_rows(tables, MergeOptions())
    assert header == ["姓名", "来源文件"]
    assert rows == [["张三", "一月.csv"]]


def test_无表头时按位置对齐并补齐宽度():
    tables = [
        ("a.csv", [["1", "2", "3"]]),
        ("b.csv", [["4", "5"]]),
    ]
    header, rows = align_rows(tables, MergeOptions(has_header=False, add_source_column=False))
    assert header == []
    assert rows == [["1", "2", "3"], ["4", "5", ""]]


def test_整行空白会被跳过():
    tables = [("a.csv", [["姓名"], ["张三"], ["  "], [""]])]
    _, rows = align_rows(tables, MergeOptions(add_source_column=False))
    assert rows == [["张三"]]


def test_参数校验():
    assert "只支持 xlsx 或 csv" in (MergeOptions(output_format="txt").validate() or "")
    assert "列名不能为空" in (MergeOptions(source_column_name=" ").validate() or "")
    assert MergeOptions().validate() is None


# ---------- 读取 ----------

def test_读取csv和xlsx(tmp_path):
    c = write_csv(tmp_path / "a.csv", [["姓名", "金额"], ["张三", "100"]])
    assert read_table(c) == [["姓名", "金额"], ["张三", "100"]]

    x = write_xlsx(tmp_path / "b.xlsx", [["姓名", "金额"], ["李四", 200]])
    assert read_table(x) == [["姓名", "金额"], ["李四", "200"]]


def test_读取GBK编码的csv(tmp_path):
    p = tmp_path / "gbk.csv"
    write_csv(p, [["姓名", "城市"], ["张三", "北京"]], encoding="gb18030")
    assert read_table(p) == [["姓名", "城市"], ["张三", "北京"]]


def test_制表符分隔的csv也能读(tmp_path):
    p = tmp_path / "tab.csv"
    p.write_text("姓名\t金额\n张三\t100\n", encoding="utf-8")
    assert read_table(p) == [["姓名", "金额"], ["张三", "100"]]


# ---------- 合并 ----------

def test_csv和xlsx混合合并成xlsx(tmp_path):
    a = write_csv(tmp_path / "一月.csv", [["姓名", "金额"], ["张三", "100"]])
    b = write_xlsx(tmp_path / "二月.xlsx", [["姓名", "金额"], ["李四", "200"]])
    target = tmp_path / "输出" / "合并.xlsx"

    opts = MergeOptions()
    report = table_job.execute(table_job.plan_merge([a, b], opts, target), opts, target)

    assert report.succeeded == 2
    wb = load_workbook(target)
    ws = wb.active
    assert [c.value for c in ws[1]] == ["姓名", "金额", "来源文件"]
    assert [c.value for c in ws[2]] == ["张三", "100", "一月.csv"]
    assert [c.value for c in ws[3]] == ["李四", "200", "二月.xlsx"]
    assert a.exists() and b.exists(), "原文件必须保留"


def test_合并输出成csv(tmp_path):
    a = write_csv(tmp_path / "a.csv", [["姓名"], ["张三"]])
    target = tmp_path / "输出" / "合并.csv"
    opts = MergeOptions(output_format="csv", add_source_column=False)
    report = table_job.execute(table_job.plan_merge([a], opts, target), opts, target)

    assert report.succeeded == 1
    assert target.read_text(encoding="utf-8-sig").splitlines() == ["姓名", "张三"]


def test_坏文件被跳过其余正常合并(tmp_path):
    a = write_csv(tmp_path / "a.csv", [["姓名"], ["张三"]])
    bad = tmp_path / "坏.xlsx"
    bad.write_bytes(b"not a real xlsx")
    target = tmp_path / "输出" / "合并.xlsx"

    opts = MergeOptions(add_source_column=False)
    report = table_job.execute(table_job.plan_merge([a, bad], opts, target), opts, target)
    assert report.succeeded == 1
    assert len(report.skipped) == 1


def test_旧版xls给出明确提示(tmp_path):
    f = tmp_path / "老表.xls"
    f.write_bytes(b"old excel")
    actions = table_job.plan_merge([f], MergeOptions(), tmp_path / "out.xlsx")
    assert actions[0].skip_reason and "另存为 .xlsx" in actions[0].skip_reason


def test_预览显示每个文件多少行多少列(tmp_path):
    a = write_csv(tmp_path / "a.csv", [["姓名", "金额"], ["张三", "100"], ["李四", "200"]])
    actions = table_job.plan_merge([a], MergeOptions(), tmp_path / "out.xlsx")
    assert "2 行数据，2 列" in actions[0].note


def test_不覆盖已存在的输出文件(tmp_path):
    a = write_csv(tmp_path / "a.csv", [["姓名"], ["张三"]])
    out = tmp_path / "输出"
    out.mkdir()
    existing = out / "合并.xlsx"
    existing.write_bytes(b"original")

    opts = MergeOptions()
    report = table_job.execute(table_job.plan_merge([a], opts, existing), opts, existing)
    assert report.succeeded == 1
    assert existing.read_bytes() == b"original"
    assert (out / "合并_1.xlsx").exists()
