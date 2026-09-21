"""所有功能页的公共安全规则 + 各自的端到端流程。

公共规则做成参数化：新加页面时会自动被这些规则覆盖，不会漏测。
"""
from pathlib import Path

import pytest

from tests.conftest import 等待线程结束


ALL_PAGE_CLASSES = "全部页面"


@pytest.fixture
def 静音弹窗(monkeypatch):
    from filebatch.ui.pages import base_page as bp

    记录 = {"info": [], "error": [], "confirm": []}
    monkeypatch.setattr(bp, "show_info", lambda p, t, m: 记录["info"].append((t, m)))
    monkeypatch.setattr(bp, "show_error", lambda p, t, m, d="": 记录["error"].append((t, m)))
    monkeypatch.setattr(bp, "confirm_dangerous", lambda *a, **k: True)
    return 记录


def 所有页面类():
    """含 Excel 工具箱里的六个子页面——它们同样要守这些公共规则。"""
    from filebatch.ui.main_window import all_job_page_classes

    return all_job_page_classes()


def 页面名(cls):
    return cls.__name__


# ---------------- 公共安全规则：每个页面都必须满足 ----------------

@pytest.mark.parametrize("page_cls", 所有页面类(), ids=页面名)
def test_每个页面执行按钮初始都是禁用的(qapp, 静音弹窗, page_cls):
    page = page_cls()
    assert not page.btn_execute.isEnabled()
    assert page.TITLE, "每个页面都要有中文标题"
    assert page.SUBTITLE, "每个页面都要有一句说明"
    page.deleteLater()


@pytest.mark.parametrize("page_cls", 所有页面类(), ids=页面名)
def test_每个页面没文件时预览都会被挡住(qapp, 静音弹窗, page_cls):
    page = page_cls()
    assert page.generate_preview() is False
    assert 静音弹窗["info"], "应该提示用户先添加文件"
    assert not page.btn_execute.isEnabled()
    page.deleteLater()


@pytest.mark.parametrize("page_cls", 所有页面类(), ids=页面名)
def test_每个页面没预览就执行都会被挡住(qapp, 静音弹窗, page_cls):
    page = page_cls()
    assert page.start_execute() is False
    page.deleteLater()


@pytest.mark.parametrize("page_cls", 所有页面类(), ids=页面名)
def test_危险开关默认关闭且警告默认隐藏(qapp, 静音弹窗, page_cls):
    page = page_cls()
    if page.chk_danger is not None:
        assert not page.chk_danger.isChecked(), "危险开关必须默认关闭"
        assert not page.lbl_danger_warn.isVisibleTo(page), "警告默认不显示"
        assert page.DANGER_CONFIRM_OK, "必须有明确的确认按钮文案"
        page.chk_danger.setChecked(True)
        assert page.lbl_danger_warn.isVisibleTo(page), "勾选后必须立刻出现警告"
    page.deleteLater()


@pytest.mark.parametrize("page_cls", 所有页面类(), ids=页面名)
def test_每个页面都能正确重置预览状态(qapp, 静音弹窗, page_cls):
    page = page_cls()
    page._invalidate_preview()
    assert page._actions == []
    assert page.preview.rowCount() == 0
    assert "预览" in page.box_preview.title()
    assert not page.btn_execute.isEnabled()
    page.deleteLater()


# ---------------- 各页面端到端 ----------------

def test_分类页把文件分到不同文件夹(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.pages.classify_page import ClassifyPage

    src = tmp_path / "杂乱"
    src.mkdir()
    (src / "照片.jpg").write_bytes(b"x")
    (src / "报表.xlsx").write_bytes(b"x")
    (src / "说明.txt").write_bytes(b"x")

    page = ClassifyPage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "整理后"

    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    assert page._last_report.succeeded == 3
    assert (tmp_path / "整理后" / "图片" / "照片.jpg").exists()
    assert (tmp_path / "整理后" / "表格" / "报表.xlsx").exists()
    assert (tmp_path / "整理后" / "文档" / "说明.txt").exists()
    assert (src / "照片.jpg").exists(), "默认复制，原文件必须保留"
    page.deleteLater()


def test_图片页改名并缩放(qapp, 静音弹窗, tmp_path):
    from PIL import Image

    from filebatch.ui.pages.image_page import ImagePage
    from filebatch.core.image_job import ResizeMode

    src = tmp_path / "图"
    src.mkdir()
    for i in range(3):
        Image.new("RGB", (1200, 800), (100, 150, 200)).save(src / f"IMG_{i}.jpg")

    page = ImagePage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "输出"
    page.ed_base.setText("产品图")
    page.cb_number.setCurrentIndex(1)
    page.sp_digits.setValue(2)
    idx = page.cb_resize.findData(ResizeMode.MAX_SIDE)
    page.cb_resize.setCurrentIndex(idx)
    page.sp_value.setValue(600)

    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    assert page._last_report.succeeded == 3
    out = tmp_path / "输出"
    assert sorted(p.name for p in out.iterdir()) == ["产品图_01.jpg", "产品图_02.jpg", "产品图_03.jpg"]
    with Image.open(out / "产品图_01.jpg") as im:
        assert im.size == (600, 400)
    page.deleteLater()


def test_图片页只接受图片文件(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.pages.image_page import ImagePage

    src = tmp_path / "混"
    src.mkdir()
    (src / "照片.jpg").write_bytes(b"x")
    (src / "说明.txt").write_bytes(b"x")
    (src / "程序.exe").write_bytes(b"x")

    page = ImagePage()
    page.file_list.add_paths_sync([src])
    assert [f.name for f in page.file_list.files()] == ["照片.jpg"], "非图片不该进清单"
    page.deleteLater()


def test_文本页替换内容(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.pages.text_page import TextPage

    src = tmp_path / "文档"
    src.mkdir()
    for i in range(2):
        (src / f"doc{i}.txt").write_text("旧公司名 出品", encoding="utf-8")

    page = TextPage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "输出"
    page.ed_find.setText("旧公司名")
    page.ed_replace.setText("新公司名")

    assert page.generate_preview() is True
    assert "将替换 1 处" in page._actions[0].note
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    assert page._last_report.succeeded == 2
    assert (tmp_path / "输出" / "doc0.txt").read_text(encoding="utf-8") == "新公司名 出品"
    assert (src / "doc0.txt").read_text(encoding="utf-8") == "旧公司名 出品", "原文件必须不变"
    page.deleteLater()


def test_文本页正则写错给中文提示(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.pages.text_page import TextPage

    f = tmp_path / "a.txt"
    f.write_text("x", encoding="utf-8")
    page = TextPage()
    page.file_list.add_paths_sync([f])
    page._output_path = tmp_path / "out"
    page.chk_regex.setChecked(True)
    page.ed_find.setText("[未闭合")

    assert page.generate_preview() is False
    assert any("正则" in m for _, m in 静音弹窗["error"])
    page.deleteLater()


def test_PDF页拆分(qapp, 静音弹窗, tmp_path):
    from pypdf import PdfReader, PdfWriter

    from filebatch.ui.pages.pdf_page import PdfPage

    src = tmp_path / "pdf"
    src.mkdir()
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=200, height=200)
    with open(src / "合同.pdf", "wb") as f:
        writer.write(f)

    page = PdfPage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    assert page._last_report.succeeded == 3
    out = tmp_path / "输出"
    assert sorted(p.name for p in out.iterdir()) == ["合同_第1页.pdf", "合同_第2页.pdf", "合同_第3页.pdf"]
    assert len(PdfReader(str(out / "合同_第1页.pdf")).pages) == 1
    page.deleteLater()


def test_PDF页合并且切模式会清掉输出位置(qapp, 静音弹窗, tmp_path):
    from pypdf import PdfReader, PdfWriter

    from filebatch.ui.pages.pdf_page import PdfPage

    src = tmp_path / "pdf"
    src.mkdir()
    for name, pages in (("a.pdf", 2), ("b.pdf", 3)):
        writer = PdfWriter()
        for _ in range(pages):
            writer.add_blank_page(width=200, height=200)
        with open(src / name, "wb") as f:
            writer.write(f)

    page = PdfPage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "先选的文件夹"
    assert page.output_is_file() is False

    page.cb_mode.setCurrentIndex(1)   # 切到合并
    assert page.output_is_file() is True
    assert page._output_path is None, "输出目标含义变了，旧的选择必须作废"

    page._output_path = tmp_path / "合并结果.pdf"
    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    assert page._last_report.succeeded == 2
    assert len(PdfReader(str(tmp_path / "合并结果.pdf")).pages) == 5
    page.deleteLater()


def test_表格页合并(qapp, 静音弹窗, tmp_path):
    import csv

    from openpyxl import load_workbook

    from filebatch.ui.pages.table_page import TablePage

    src = tmp_path / "表"
    src.mkdir()
    for name, rows in (
        ("一月.csv", [["姓名", "金额"], ["张三", "100"]]),
        ("二月.csv", [["金额", "姓名"], ["200", "李四"]]),   # 列顺序故意反过来
    ):
        with open(src / name, "w", encoding="utf-8", newline="") as f:
            csv.writer(f).writerows(rows)

    page = TablePage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "合并.xlsx"

    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    assert page._last_report.succeeded == 2
    ws = load_workbook(tmp_path / "合并.xlsx").active
    assert [c.value for c in ws[1]] == ["姓名", "金额", "来源文件"]
    assert [c.value for c in ws[2]] == ["张三", "100", "一月.csv"]
    assert [c.value for c in ws[3]] == ["李四", "200", "二月.csv"], "列顺序不同也要按列名对齐"
    page.deleteLater()


def test_表格页只接受表格文件(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.pages.table_page import TablePage

    src = tmp_path / "混"
    src.mkdir()
    (src / "数据.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    (src / "照片.jpg").write_bytes(b"x")

    page = TablePage()
    page.file_list.add_paths_sync([src])
    assert [f.name for f in page.file_list.files()] == ["数据.csv"]
    page.deleteLater()


def test_分类页开启移动后原文件会被拿走(qapp, 静音弹窗, tmp_path):
    """危险路径也要有测试：确认"移动"真的是移动，而不是偷偷复制。"""
    from filebatch.ui.pages.classify_page import ClassifyPage

    src = tmp_path / "源"
    src.mkdir()
    f = src / "照片.jpg"
    f.write_bytes(b"x")

    page = ClassifyPage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "整理后"
    page.chk_danger.setChecked(True)

    assert page.generate_preview() is True
    assert page.start_execute() is True     # 静音弹窗里 confirm 固定返回 True
    等待线程结束(qapp, page._worker)

    assert page._last_report.succeeded == 1
    assert (tmp_path / "整理后" / "图片" / "照片.jpg").exists()
    assert not f.exists(), "移动模式下原文件应该已经不在原位置"
    page.deleteLater()
