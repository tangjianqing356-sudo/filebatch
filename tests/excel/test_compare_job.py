"""两表对比。"""
from openpyxl import load_workbook

from filebatch.core.excel import compare_job
from filebatch.core.excel.compare_job import (
    SHEET_DIFF,
    SHEET_ONLY_A,
    SHEET_ONLY_B,
    SHEET_SUMMARY,
    CompareOptions,
    compare_tables,
)
from filebatch.core.excel.workbook import TableData


def _表(rows, header=("工号", "姓名", "部门")):
    return TableData(header=list(header), rows=rows)


def test_找出仅A表和仅B表有的():
    a = _表([["1", "张三", "研发"], ["2", "李四", "市场"]])
    b = _表([["2", "李四", "市场"], ["3", "王五", "销售"]])
    r = compare_tables(a, b, CompareOptions(key_columns=["工号"]))

    assert [row[0] for row in r.only_a] == ["1"]
    assert [row[0] for row in r.only_b] == ["3"]
    assert r.same_count == 1
    assert r.diff_rows == []


def test_找出内容不同的():
    a = _表([["1", "张三", "研发"]])
    b = _表([["1", "张三", "市场"]])
    r = compare_tables(a, b, CompareOptions(key_columns=["工号"]))

    assert len(r.diff_rows) == 1
    行 = r.diff_rows[0]
    assert 行[0] == "1" and 行[1] == "部门" and 行[2] == "研发" and 行[3] == "市场"
    assert r.same_count == 0


def test_一行有多处不同会拆成多条():
    a = _表([["1", "张三", "研发"]])
    b = _表([["1", "张四", "市场"]])
    r = compare_tables(a, b, CompareOptions(key_columns=["工号"]))
    assert len(r.diff_rows) == 2
    assert {行[1] for 行 in r.diff_rows} == {"姓名", "部门"}


def test_多列联合做关键列():
    a = _表([["1", "张三", "研发"], ["1", "李四", "市场"]])
    b = _表([["1", "张三", "销售"], ["1", "李四", "市场"]])
    r = compare_tables(a, b, CompareOptions(key_columns=["工号", "姓名"]))
    assert len(r.diff_rows) == 1
    assert r.same_count == 1


def test_只比指定的列():
    a = _表([["1", "张三", "研发"]])
    b = _表([["1", "张四", "市场"]])
    r = compare_tables(a, b, CompareOptions(key_columns=["工号"], compare_columns=["部门"]))
    assert len(r.diff_rows) == 1
    assert r.diff_rows[0][1] == "部门", "没要求比姓名就不该报姓名"


def test_关键列重复时保留第一条并计数():
    """静默丢掉重复行会让用户以为数据没了，必须显式报出来。"""
    a = _表([["1", "张三", "研发"], ["1", "重复的", "市场"]])
    b = _表([["1", "张三", "研发"]])
    r = compare_tables(a, b, CompareOptions(key_columns=["工号"]))
    assert r.duplicate_keys_a == 1
    assert r.same_count == 1


def test_忽略大小写和空格():
    a = _表([["1", "  张三 ", "研发"]])
    b = _表([["1", "张三", "研发"]])
    r = compare_tables(a, b, CompareOptions(key_columns=["工号"], trim_spaces=True))
    assert r.same_count == 1 and not r.diff_rows


def test_空行不参与对比():
    a = _表([["1", "张三", "研发"], ["", "", ""]])
    b = _表([["1", "张三", "研发"]])
    r = compare_tables(a, b, CompareOptions(key_columns=["工号"]))
    assert not r.only_a and r.same_count == 1


def test_参数校验():
    assert "关键列" in (CompareOptions().validate() or "")
    assert "表头" in (CompareOptions(key_columns=["A"], has_header=False).validate() or "")
    assert CompareOptions(key_columns=["工号"]).validate() is None


# ---------------- 计划与执行 ----------------

def test_必须正好两个文件(写xlsx, tmp_path):
    a = 写xlsx(tmp_path / "a.xlsx", [["工号"], ["1"]])
    opts = CompareOptions(key_columns=["工号"])

    for files in ([a], [a, a, a]):
        actions = compare_job.plan(files, opts, tmp_path / "out.xlsx")
        assert not actions[0].will_run
        assert "正好 2 个文件" in actions[0].skip_reason


def test_关键列不存在时指明是哪张表(写xlsx, tmp_path):
    a = 写xlsx(tmp_path / "a.xlsx", [["工号", "姓名"], ["1", "张三"]])
    b = 写xlsx(tmp_path / "b.xlsx", [["编号", "姓名"], ["1", "张三"]])
    actions = compare_job.plan([a, b], CompareOptions(key_columns=["工号"]), tmp_path / "out.xlsx")
    assert not actions[0].will_run
    assert "B 表里找不到关键列" in actions[0].skip_reason


def test_完整对比并输出四张工作表(写xlsx, tmp_path):
    a = 写xlsx(tmp_path / "员工A.xlsx",
               [["工号", "姓名", "部门"], ["1", "张三", "研发"], ["2", "李四", "市场"]])
    b = 写xlsx(tmp_path / "员工B.xlsx",
               [["工号", "姓名", "部门"], ["2", "李四", "销售"], ["3", "王五", "运营"]])
    out = tmp_path / "对比结果.xlsx"
    opts = CompareOptions(key_columns=["工号"])

    actions = compare_job.plan([a, b], opts, out)
    assert actions[0].will_run
    assert not out.exists(), "计划阶段不写文件"

    report = compare_job.execute(actions, opts, out)
    assert report.succeeded == 1 and report.failed == 0

    wb = load_workbook(out)
    assert wb.sheetnames == [SHEET_SUMMARY, SHEET_ONLY_A, SHEET_ONLY_B, SHEET_DIFF]
    assert [c.value for c in wb[SHEET_ONLY_A][2]][0] == "1"
    assert [c.value for c in wb[SHEET_ONLY_B][2]][0] == "3"
    assert wb[SHEET_DIFF].max_row == 2, "一行表头 + 一条差异"
    assert a.exists() and b.exists(), "原文件必须保留"


def test_不覆盖已有的输出文件(写xlsx, tmp_path):
    a = 写xlsx(tmp_path / "a.xlsx", [["工号"], ["1"]])
    b = 写xlsx(tmp_path / "b.xlsx", [["工号"], ["2"]])
    out = tmp_path / "结果.xlsx"
    out.write_bytes(b"original")

    opts = CompareOptions(key_columns=["工号"])
    report = compare_job.execute(compare_job.plan([a, b], opts, out), opts, out)

    assert report.succeeded == 1
    assert out.read_bytes() == b"original", "已有文件绝不能被覆盖"
    assert (tmp_path / "结果_1.xlsx").exists()
    assert report.items[0].target.name == "结果_1.xlsx", "结果里要回填真正写入的文件名"
