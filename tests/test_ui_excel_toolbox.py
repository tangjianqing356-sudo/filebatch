"""Excel 工具箱的界面层。

公共安全规则已经在 test_ui_all_pages.py 里参数化覆盖到每个子页面了，
这里测的是工具箱自己的东西：懒加载、收尾转发、后台读列名，以及各子功能走通。
"""
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import 等待线程结束


@pytest.fixture
def 静音弹窗(monkeypatch):
    from filebatch.ui.pages import base_page as bp

    记录 = {"info": [], "error": []}
    monkeypatch.setattr(bp, "show_info", lambda p, t, m: 记录["info"].append((t, m)))
    monkeypatch.setattr(bp, "show_error", lambda p, t, m, d="": 记录["error"].append((t, m)))
    monkeypatch.setattr(bp, "confirm_dangerous", lambda *a, **k: True)
    return 记录


@pytest.fixture
def 写xlsx():
    def _写(path: Path, rows: list[list]) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        for r in rows:
            ws.append(r)
        wb.save(path)
        return path
    return _写


def 等列名(qapp, page):
    """列名是后台线程读的，测试里必须等它回来。"""
    worker = page._inspect_worker
    assert worker is not None, "添加文件后应该起一个读表头的线程"
    等待线程结束(qapp, worker)


# ---------------- 工具箱本身 ----------------

def test_工具箱子页面是按需创建的(qapp, 静音弹窗):
    from filebatch.ui.excel_toolbox.toolbox_page import SUB_FEATURES, ExcelToolboxPage

    tb = ExcelToolboxPage()
    assert tb.tabs.count() == len(SUB_FEATURES)
    assert len(tb.created_pages()) == 1, "只有当前这一页该被建出来"

    tb.tabs.setCurrentIndex(3)
    assert len(tb.created_pages()) == 2, "切过去才建第二页"
    tb.deleteLater()


def test_工具箱切标签不会串页(qapp, 静音弹窗):
    """主窗口那边踩过一次"建页面把索引挤歪"的坑，这里锁住。"""
    from filebatch.ui.excel_toolbox.toolbox_page import SUB_FEATURES, ExcelToolboxPage

    tb = ExcelToolboxPage()
    for i, (名称, cls) in enumerate(SUB_FEATURES):
        tb.tabs.setCurrentIndex(i)
        assert tb.tabs.currentIndex() == i
        assert isinstance(tb.sub_page(i), cls), f"第 {i} 个标签 {名称} 拿到的页面不对"
    tb.deleteLater()


def test_工具箱把忙碌状态转发给子页面(qapp, 静音弹窗, 写xlsx, tmp_path):
    """关窗保护靠这个：子页面在跑任务，工具箱必须说自己忙。"""
    from filebatch.ui.excel_toolbox.toolbox_page import ExcelToolboxPage

    tb = ExcelToolboxPage()
    assert tb.is_busy() is False
    assert tb.active_workers() == []

    page = tb.sub_page(3)      # 数据清洗
    src = 写xlsx(tmp_path / "源" / "a.xlsx", [["A"], ["  1  "], ["", ""]])
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "输出"
    assert page.generate_preview() is True
    assert page.start_execute() is True

    # 任务刚起来，worker 还在，工具箱必须能看见
    assert tb.is_busy() == page.is_busy()
    等待线程结束(qapp, page._worker)
    assert tb.is_busy() is False
    tb.deleteLater()


def test_工具箱的停止会转发到所有已建子页面(qapp, 静音弹窗):
    from filebatch.ui.excel_toolbox.toolbox_page import ExcelToolboxPage

    tb = ExcelToolboxPage()
    for i in range(tb.tabs.count()):
        tb.sub_page(i)
    调用 = []
    for p in tb.created_pages():
        p.stop_background_work = lambda p=p: 调用.append(type(p).__name__)

    tb.stop_background_work()
    assert len(调用) == 6, f"六个子页面都该被叫到：{调用}"
    tb.deleteLater()


def test_主窗口里工具箱占一个入口(qapp, 静音弹窗):
    from filebatch.ui.excel_toolbox.toolbox_page import ExcelToolboxPage
    from filebatch.ui.main_window import EXCEL_TOOLBOX_INDEX, MainWindow

    win = MainWindow()
    assert win.nav.item(EXCEL_TOOLBOX_INDEX).text() == "Excel 工具箱"
    win.nav.setCurrentRow(EXCEL_TOOLBOX_INDEX)
    assert win.stack.currentIndex() == EXCEL_TOOLBOX_INDEX, "切过去不能是空白页"
    assert isinstance(win.page_at(EXCEL_TOOLBOX_INDEX), ExcelToolboxPage)
    win.close()


def test_合并功能还在工具箱第一个标签里(qapp, 静音弹窗):
    """原有功能不能被这轮改动弄丢。"""
    from filebatch.ui.excel_toolbox.toolbox_page import ExcelToolboxPage
    from filebatch.ui.pages.table_page import TablePage

    tb = ExcelToolboxPage()
    assert tb.tabs.tabText(0) == "合并"
    assert isinstance(tb.sub_page(0), TablePage)
    tb.deleteLater()


# ---------------- 列名是后台读的 ----------------

def test_列名读出来会填进选择控件(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = 写xlsx(tmp_path / "源" / "a.xlsx", [["姓名", "手机", "城市"], ["张三", "1", "北京"]])
    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)

    assert page.list_keys.count() == 3
    assert [page.list_keys.item(i).text() for i in range(3)] == ["姓名", "手机", "城市"]
    assert page.list_keys.checked() == [], "默认一个都不勾 = 整行比较"
    page.deleteLater()


def test_多个文件只给共有列(qapp, 静音弹窗, 写xlsx, tmp_path):
    """按只有部分文件才有的列去重，剩下的文件必然失败——干脆别让用户选到。"""
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = tmp_path / "源"
    写xlsx(src / "a.xlsx", [["工号", "姓名", "部门"], ["1", "张三", "研发"]])
    写xlsx(src / "b.xlsx", [["工号", "姓名", "城市"], ["1", "张三", "北京"]])

    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)

    assert [page.list_keys.item(i).text() for i in range(page.list_keys.count())] == ["工号", "姓名"]
    page.deleteLater()


def test_读不了的文件在界面上说清楚(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = tmp_path / "源"
    src.mkdir()
    坏 = src / "老表.xlsx"
    坏.write_bytes(b"\xd0\xcf\x11\xe0" + b"x" * 100)     # 伪装成 xlsx 的 OLE2

    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)

    提示 = page.lbl_columns.text()
    assert "老表.xlsx" in 提示 and "读不了" in 提示, 提示
    page.deleteLater()


def test_勾好的列在重新读表头后还在(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = tmp_path / "源"
    写xlsx(src / "a.xlsx", [["姓名", "手机"], ["张三", "1"]])
    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page.list_keys.set_checked(["姓名"])

    写xlsx(src / "b.xlsx", [["姓名", "手机"], ["李四", "2"]])
    page.file_list.add_paths_sync([src / "b.xlsx"])
    等列名(qapp, page)

    assert page.list_keys.checked() == ["姓名"], "重新读表头不该把用户的选择清掉"
    page.deleteLater()


# ---------------- 各子功能走通 ----------------

def test_去重页(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = 写xlsx(tmp_path / "源" / "客户.xlsx", [
        ["姓名", "手机", "城市"],
        ["张三", "138", "北京"],
        ["李四", "139", "上海"],
        ["张三", "138", "广州"],
    ])
    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page.list_keys.set_checked(["姓名", "手机"])
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True
    assert "重复" in page._actions[0].note, page._actions[0].note
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    assert page._last_report.succeeded == 1
    ws = load_workbook(tmp_path / "输出" / "客户_去重.xlsx").active
    assert ws.max_row == 3
    assert ws["C2"].value == "北京", "保留第一条"
    assert src.exists(), "原文件必须保留"
    page.deleteLater()


def test_去重页可以保留最后一条(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = 写xlsx(tmp_path / "源" / "客户.xlsx", [
        ["姓名", "城市"], ["张三", "北京"], ["张三", "广州"],
    ])
    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page.list_keys.set_checked(["姓名"])
    page.cb_keep.setCurrentIndex(1)
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    ws = load_workbook(tmp_path / "输出" / "客户_去重.xlsx").active
    assert ws["B2"].value == "广州"
    page.deleteLater()


def test_对比页(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.compare_page import ComparePage

    src = tmp_path / "源"
    写xlsx(src / "1_上月.xlsx", [["工号", "金额"], ["001", "100"], ["002", "200"]])
    写xlsx(src / "2_本月.xlsx", [["工号", "金额"], ["001", "100"], ["003", "300"]])

    page = ComparePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    assert page.selected_pair()[0].name == "1_上月.xlsx", "默认 A 是清单里第一个"
    page.list_keys.set_checked(["工号"])
    page._output_path = tmp_path / "对比结果.xlsx"

    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    wb = load_workbook(tmp_path / "对比结果.xlsx")
    assert wb.sheetnames == ["汇总", "仅A表有", "仅B表有", "内容不同"], "汇总放最前，打开就看到结论"
    assert wb["仅A表有"].max_row == 2
    assert wb["仅B表有"].max_row == 2
    page.deleteLater()


# ---------------- 两表对比的 A / B 选择 ----------------

def 建对比页(qapp, 写xlsx, tmp_path, 名字: list[str]):
    from filebatch.ui.excel_toolbox.compare_page import ComparePage

    src = tmp_path / "源"
    for n in 名字:
        写xlsx(src / n, [["工号", "金额"], ["001", "100"]])
    page = ComparePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    return page, src


def test_对比页默认选前两个文件(qapp, 静音弹窗, 写xlsx, tmp_path):
    page, _ = 建对比页(qapp, 写xlsx, tmp_path, ["1_上月.xlsx", "2_本月.xlsx", "3_下月.xlsx"])

    a, b = page.selected_pair()
    assert a.name == "1_上月.xlsx"
    assert b.name == "2_本月.xlsx"
    assert "A = 1_上月.xlsx" in page.lbl_ab.text()
    assert "B = 2_本月.xlsx" in page.lbl_ab.text()
    page.deleteLater()


def test_对比页可以从一堆文件里挑两个(qapp, 静音弹窗, 写xlsx, tmp_path):
    """清单里放三个文件不再是错误——挑哪两个由用户说了算。"""
    page, _ = 建对比页(qapp, 写xlsx, tmp_path, ["a.xlsx", "b.xlsx", "c.xlsx"])

    page.cb_a.setCurrentIndex(page.cb_a.findData(str(page._paths()[0])))
    page.cb_b.setCurrentIndex(page.cb_b.findData(str(page._paths()[2])))
    等列名(qapp, page)

    a, b = page.selected_pair()
    assert (a.name, b.name) == ("a.xlsx", "c.xlsx")
    assert page.validate_options() is None or "关键列" in page.validate_options()
    page.deleteLater()


def test_对比页不允许两边选同一个文件(qapp, 静音弹窗, 写xlsx, tmp_path):
    page, _ = 建对比页(qapp, 写xlsx, tmp_path, ["a.xlsx", "b.xlsx", "c.xlsx"])

    # 把 A 挪到 B 现在选的那个文件上
    b_path = page.selected_pair()[1]
    page.cb_a.setCurrentIndex(page.cb_a.findData(str(b_path)))

    a, b = page.selected_pair()
    assert a == b_path
    assert b is not None and b != a, "撞车时另一边必须自动让开"
    page.deleteLater()


def test_对比页只有一个文件时说清楚(qapp, 静音弹窗, 写xlsx, tmp_path):
    page, _ = 建对比页(qapp, 写xlsx, tmp_path, ["只有这一个.xlsx"])

    a, b = page.selected_pair()
    assert a is not None and b is None
    assert "至少放 2 个" in page.lbl_ab.text(), page.lbl_ab.text()
    assert "A 表和 B 表" in (page.validate_options() or "")
    page.deleteLater()


def test_对比页交换AB(qapp, 静音弹窗, 写xlsx, tmp_path):
    page, _ = 建对比页(qapp, 写xlsx, tmp_path, ["a.xlsx", "b.xlsx"])
    原a, 原b = page.selected_pair()

    page.btn_swap.click()

    新a, 新b = page.selected_pair()
    assert (新a, 新b) == (原b, 原a)
    page.deleteLater()


def test_对比页文件被删掉后选择状态会同步(qapp, 静音弹窗, 写xlsx, tmp_path):
    """清单里删掉正在用的文件，下拉不能还指着一个不存在的路径。"""
    page, src = 建对比页(qapp, 写xlsx, tmp_path, ["a.xlsx", "b.xlsx", "c.xlsx"])
    page.cb_b.setCurrentIndex(page.cb_b.findData(str(src / "c.xlsx")))
    assert page.selected_pair()[1].name == "c.xlsx"

    # 模拟用户在清单里把 c.xlsx 移除
    page.file_list.clear()
    page.file_list.add_paths_sync([src / "a.xlsx", src / "b.xlsx"])
    qapp.processEvents()

    a, b = page.selected_pair()
    剩下 = set(page.file_list.files())
    assert a in 剩下 and b in 剩下, "选中的必须还在清单里"
    assert a != b
    assert page.cb_a.count() == 2 and page.cb_b.count() == 2
    page.deleteLater()


def test_对比页清空文件后不会留下悬空选择(qapp, 静音弹窗, 写xlsx, tmp_path):
    page, _ = 建对比页(qapp, 写xlsx, tmp_path, ["a.xlsx", "b.xlsx"])
    page.file_list.clear()
    qapp.processEvents()

    assert page.selected_pair() == (None, None)
    assert page.cb_a.count() == 0
    assert "A 表和 B 表" in (page.validate_options() or "")
    page.deleteLater()


def test_对比页换了AB会重读列名(qapp, 静音弹窗, 写xlsx, tmp_path):
    """A/B 换成表头不同的文件，列下拉必须跟着变，不能还是上一对的列。"""
    from filebatch.ui.excel_toolbox.compare_page import ComparePage

    src = tmp_path / "源"
    写xlsx(src / "a.xlsx", [["工号", "姓名"], ["1", "张三"]])
    写xlsx(src / "b.xlsx", [["工号", "姓名"], ["1", "张三"]])
    写xlsx(src / "c.xlsx", [["单号", "金额"], ["X1", "9"]])

    page = ComparePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    assert [page.list_keys.item(i).text() for i in range(page.list_keys.count())] == ["工号", "姓名"]

    page.cb_b.setCurrentIndex(page.cb_b.findData(str(src / "c.xlsx")))
    等列名(qapp, page)

    assert page.list_keys.count() == 0, "a 和 c 没有共同的列"
    page.deleteLater()


def test_清洗页(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.clean_page import CleanPage

    src = 写xlsx(tmp_path / "源" / "名单.xlsx", [
        ["姓名", "备注", "编号"],
        ["  张三  ", "多个   空格", "001"],
        ["", "", ""],
        ["李四", "换行", "0012"],
    ])
    page = CleanPage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    ws = load_workbook(tmp_path / "输出" / "名单_清洗.xlsx").active
    assert ws.max_row == 3, "空行该被删掉"
    assert ws["A2"].value == "张三"
    assert ws["B2"].value == "多个 空格"
    assert ws["C2"].value == "001", "以 0 开头的编号不能被当成数字"
    assert ws["C3"].value == "0012"
    page.deleteLater()


def test_清洗页至少要勾一项(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.clean_page import CleanPage

    src = 写xlsx(tmp_path / "源" / "a.xlsx", [["A"], ["1"]])
    page = CleanPage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "输出"
    for c in (page.chk_blank_rows, page.chk_trim, page.chk_collapse, page.chk_newline):
        c.setChecked(False)

    assert page.generate_preview() is False
    assert any("至少" in m for _, m in 静音弹窗["error"]), 静音弹窗["error"]
    page.deleteLater()


def test_拆表页(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.split_page import SplitPage

    src = 写xlsx(tmp_path / "源" / "花名册.xlsx", [
        ["姓名", "部门"],
        ["张三", "研发"],
        ["李四", "销售"],
        ["王五", "研发"],
    ])
    page = SplitPage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page.cb_column.setCurrentText("部门")
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True
    assert len([a for a in page._actions if a.will_run]) == 2
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    out = tmp_path / "输出"
    assert sorted(p.name for p in out.iterdir()) == ["花名册_研发.xlsx", "花名册_销售.xlsx"]
    assert load_workbook(out / "花名册_研发.xlsx").active.max_row == 3
    page.deleteLater()


def test_拆表页遇到文件名非法字符不会炸(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.split_page import SplitPage

    src = 写xlsx(tmp_path / "源" / "表.xlsx", [
        ["名", "部门"], ["a", "行政/后勤"], ["b", "研发:一组"],
    ])
    page = SplitPage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page.cb_column.setCurrentText("部门")
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    assert page._last_report.failed == 0, page._last_report.summary()
    names = [p.name for p in (tmp_path / "输出").iterdir()]
    assert len(names) == 2, names
    assert not any("/" in n or ":" in n for n in names), names
    page.deleteLater()


def test_拆表页拆出太多会被拦住(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.split_page import SplitPage

    src = 写xlsx(tmp_path / "源" / "订单.xlsx",
                [["订单号", "金额"]] + [[f"SN{i:04d}", "1"] for i in range(50)])
    page = SplitPage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page.cb_column.setCurrentText("订单号")
    page.sp_max.setValue(10)
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is False, "全都超上限时不该让用户点执行"
    assert not list((tmp_path / "输出").iterdir()) if (tmp_path / "输出").exists() else True
    page.deleteLater()


def 写带字体(path, rows, 字体表=None):
    from openpyxl.styles import Font

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r in rows:
        ws.append(r)
    for 坐标, (名, 号) in (字体表 or {}).items():
        ws[坐标].font = Font(name=名, size=号)
    wb.save(path)
    return path


def 跑格式页(qapp, page, tmp_path, 源名="报表"):
    page._output_path = tmp_path / "输出"
    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)
    assert page._last_report.failed == 0, page._last_report.summary()
    return load_workbook(tmp_path / "输出" / f"{源名}_格式化.xlsx").active


def test_格式页(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    src = 写xlsx(tmp_path / "源" / "报表.xlsx", [["姓名", "金额"], ["张三", 100]])
    page = FormatPage()
    page.file_list.add_paths_sync([src])

    ws = 跑格式页(qapp, page, tmp_path)
    assert ws["A1"].font.bold is True
    assert ws.freeze_panes == "A2"
    assert ws["A2"].border.left.style == "thin"
    assert src.exists(), "原文件必须保留"
    page.deleteLater()


# ---------------- 格式页的字体：默认不改 ----------------

def test_格式页默认是保留原字体(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    page = FormatPage()
    assert page.cb_font.currentIndex() == 0
    assert "保留原字体" in page.cb_font.currentText()
    assert page.selected_font() is None
    assert page.selected_size() is None

    opts = page.current_options()
    assert opts.font_name is None and opts.font_size is None
    page.deleteLater()


def test_格式页默认执行后原字体不变(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    src = 写带字体(
        tmp_path / "源" / "报表.xlsx",
        [["姓名", "金额"], ["张三", 100], ["李四", 200]],
        {"A1": ("Times New Roman", 13), "A2": ("宋体", 14), "A3": ("Arial", 9)},
    )
    page = FormatPage()
    page.file_list.add_paths_sync([src])

    ws = 跑格式页(qapp, page, tmp_path)
    assert (ws["A2"].font.name, ws["A2"].font.size) == ("宋体", 14)
    assert (ws["A3"].font.name, ws["A3"].font.size) == ("Arial", 9)
    assert ws["A1"].font.name == "Times New Roman"
    assert ws["A1"].font.bold is True, "加粗是用户要的，字体不是"
    page.deleteLater()


def test_格式页选了字体才统一(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    src = 写带字体(
        tmp_path / "源" / "报表.xlsx",
        [["姓名"], ["张三"], ["李四"]],
        {"A2": ("宋体", 14), "A3": ("Arial", 9)},
    )
    page = FormatPage()
    page.file_list.add_paths_sync([src])
    page.cb_font.setCurrentText("微软雅黑")
    assert page.selected_font() == "微软雅黑"

    ws = 跑格式页(qapp, page, tmp_path)
    assert ws["A1"].font.name == "微软雅黑"
    assert ws["A2"].font.name == "微软雅黑"
    assert ws["A3"].font.name == "微软雅黑"
    assert ws["A2"].font.size == 14, "没动字号就该保持原样"
    page.deleteLater()


def test_格式页可以只改字号(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    src = 写带字体(tmp_path / "源" / "报表.xlsx", [["A"], ["1"]], {"A2": ("宋体", 14)})
    page = FormatPage()
    page.file_list.add_paths_sync([src])
    page.sp_size.setValue(11)
    assert page.selected_size() == 11 and page.selected_font() is None

    ws = 跑格式页(qapp, page, tmp_path)
    assert ws["A2"].font.name == "宋体"
    assert ws["A2"].font.size == 11
    page.deleteLater()


def test_格式页字体下拉里有保留项和常用字体(qapp, 静音弹窗):
    from filebatch.core.excel.format_job import FONT_CHOICES
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    page = FormatPage()
    项 = [page.cb_font.itemText(i) for i in range(page.cb_font.count())]
    assert "保留原字体" in 项[0]
    for 名 in ("微软雅黑", "宋体", "Arial"):
        assert 名 in 项, f"{名} 应该在常用字体里：{项}"
    assert list(FONT_CHOICES) == 项[1:]
    page.deleteLater()


def test_格式页可以手打列表里没有的字体(qapp, 静音弹窗):
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    page = FormatPage()
    assert page.cb_font.isEditable(), "得允许用户打别的字体名"
    page.cb_font.setCurrentText("方正舒体")
    assert page.selected_font() == "方正舒体"
    page.deleteLater()


def test_格式页多工作表不会被弄丢(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    src = tmp_path / "源" / "多表.xlsx"
    src.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "一月"
    ws.append(["姓名"]); ws.append(["张三"])
    w2 = wb.create_sheet("二月")
    w2.append(["姓名"]); w2.append(["李四"])
    wb.save(src)

    page = FormatPage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "输出"
    assert page.generate_preview() is True
    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)

    得到 = load_workbook(tmp_path / "输出" / "多表_格式化.xlsx")
    assert 得到.sheetnames == ["一月", "二月"]
    page.deleteLater()


def test_格式页只收xlsx(qapp, 静音弹窗, tmp_path):
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    src = tmp_path / "源"
    src.mkdir()
    (src / "数据.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    (src / "表.xlsx").write_bytes(b"PK\x03\x04")

    page = FormatPage()
    page.file_list.add_paths_sync([src])
    assert [f.name for f in page.file_list.files()] == ["表.xlsx"], "CSV 不该进清单"
    page.deleteLater()


def test_格式页对图表的说法要和实际行为一致(qapp, 静音弹窗, tmp_path):
    from openpyxl.chart import BarChart, Reference

    from filebatch.ui.excel_toolbox.format_page import FormatPage

    src = tmp_path / "源"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    for r in [["月份", "销量"], ["一月", 10], ["二月", 20]]:
        ws.append(r)
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    ws.add_chart(chart, "D2")
    wb.save(src / "带图表.xlsx")

    page = FormatPage()
    page.file_list.add_paths_sync([src])
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True, "只提示，不阻止"
    note = page._actions[0].note
    assert "图表" in note and "已验证可保留" in note, note
    assert "兼容性差异" in note, "不能打包票"
    assert "一定" not in note and "保证" not in note, f"别做绝对承诺：{note}"

    assert page.start_execute() is True
    等待线程结束(qapp, page._worker)
    ws = load_workbook(tmp_path / "输出" / "带图表_格式化.xlsx").active
    assert ws._charts, "说了验收样例能保留，这个样例就真得保留"
    page.deleteLater()


def test_子页面各自独立互不干扰(qapp, 静音弹窗, 写xlsx, tmp_path):
    """两个标签页各自有文件清单和输出位置，不能串。"""
    from filebatch.ui.excel_toolbox.toolbox_page import ExcelToolboxPage

    tb = ExcelToolboxPage()
    去重 = tb.sub_page(1)
    清洗 = tb.sub_page(3)

    a = 写xlsx(tmp_path / "源" / "a.xlsx", [["A"], ["1"]])
    去重.file_list.add_paths_sync([a])
    去重._output_path = tmp_path / "去重输出"

    assert 清洗.file_list.files() == []
    assert 清洗._output_path is None
    tb.deleteLater()


# ---------------- 预览前的全量表头体检 ----------------
#
# 列下拉只读前 20 个文件，所以用户可能选到一个"后面的文件没有"的列。
# 预览阶段必须把全部文件体检一遍，问题在执行前就摆出来。

def test_预览会检查超出前20个的文件(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.columns import MAX_INSPECT_FILES
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = tmp_path / "源"
    # 前 20 个有「手机」列，第 21 个没有
    for i in range(MAX_INSPECT_FILES):
        写xlsx(src / f"{i:02d}_正常.xlsx",
              [["姓名", "手机"], ["张三", "138"], ["张三", "138"]])
    写xlsx(src / "99_缺列.xlsx", [["姓名"], ["张三"], ["张三"]])

    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    # 下拉里看得到「手机」，因为前 20 个都有
    assert "手机" in [page.list_keys.item(i).text() for i in range(page.list_keys.count())]
    page.list_keys.set_checked(["手机"])
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True

    报告 = page._header_report
    assert 报告 is not None, "预览必须跑过全量体检"
    assert 报告.total == MAX_INSPECT_FILES + 1, "要检查全部文件，不是只检查前 20 个"
    assert [p.name for p in 报告.blocked] == ["99_缺列.xlsx"]

    # 问题文件在日志里写明了
    日志 = page.log.toPlainText()
    assert "99_缺列.xlsx" in 日志 and "缺少「手机」" in 日志, 日志

    # 也在预览表里以跳过项的形式出现
    缺列项 = [a for a in page._actions if a.source.name == "99_缺列.xlsx"]
    assert len(缺列项) == 1
    assert not 缺列项[0].will_run
    assert "手机" in 缺列项[0].skip_reason
    page.deleteLater()


def test_表头不一致的文件会在预览里被标出来(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = tmp_path / "源"
    写xlsx(src / "a.xlsx", [["姓名", "手机", "部门"], ["张三", "138", "研发"], ["张三", "138", "研发"]])
    写xlsx(src / "b.xlsx", [["姓名", "手机", "城市"], ["李四", "139", "北京"], ["李四", "139", "北京"]])

    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page.list_keys.set_checked(["姓名"])
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True

    报告 = page._header_report
    assert [p.name for p in 报告.mismatched] == ["b.xlsx"]
    assert 报告.blocked == [], "表头不一样不该拦住处理"

    b项 = [a for a in page._actions if a.source.name == "b.xlsx"][0]
    assert b项.will_run, "需要的列都在，照常处理"
    assert "表头和其它文件不一致" in b项.note, b项.note
    assert "b.xlsx" in page.log.toPlainText()
    page.deleteLater()


def test_读不了的文件在预览阶段就被点名(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = tmp_path / "源"
    写xlsx(src / "好.xlsx", [["姓名"], ["张三"], ["张三"]])
    (src / "老表.xlsx").write_bytes(b"\xd0\xcf\x11\xe0" + b"x" * 100)

    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page.list_keys.set_checked(["姓名"])
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True
    assert [p.name for p in page._header_report.blocked] == ["老表.xlsx"]
    assert "老表.xlsx" in page.log.toPlainText()
    page.deleteLater()


def test_体检结果会随着换列而作废(qapp, 静音弹窗, 写xlsx, tmp_path):
    """换了关键列，上一次按旧列算出来的体检结论就不作数了。"""
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = tmp_path / "源"
    写xlsx(src / "a.xlsx", [["姓名", "手机"], ["张三", "138"], ["张三", "138"]])
    写xlsx(src / "b.xlsx", [["姓名"], ["李四"], ["李四"]])

    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page._output_path = tmp_path / "输出"

    page.list_keys.set_checked(["姓名"])
    assert page.generate_preview() is True
    assert page._header_report.blocked == [], "按姓名去重，两个文件都行"

    # 模拟"前 20 个文件里有手机列所以被列进下拉"的情况
    page.list_keys.set_columns(["姓名", "手机"])
    page.list_keys.set_checked(["手机"])
    assert page.generate_preview() is True
    assert [p.name for p in page._header_report.blocked] == ["b.xlsx"], "换成手机后 b 就缺列了"
    page.deleteLater()


def test_体检结果会随着换文件而作废(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = tmp_path / "源"
    写xlsx(src / "a.xlsx", [["姓名"], ["张三"], ["张三"]])

    page = DedupePage()
    page.file_list.add_paths_sync([src / "a.xlsx"])
    等列名(qapp, page)
    page.list_keys.set_checked(["姓名"])
    page._output_path = tmp_path / "输出"
    assert page.generate_preview() is True
    assert page._header_report.total == 1

    写xlsx(src / "b.xlsx", [["单号"], ["X1"], ["X1"]])
    page.file_list.add_paths_sync([src / "b.xlsx"])
    等列名(qapp, page)
    assert page._header_report is None, "加了文件，旧体检必须作废"

    assert page.generate_preview() is True
    assert page._header_report.total == 2
    page.deleteLater()


def test_拆表页也会全量体检(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.split_page import SplitPage

    src = tmp_path / "源"
    写xlsx(src / "a.xlsx", [["姓名", "部门"], ["张三", "研发"]])
    写xlsx(src / "b.xlsx", [["姓名", "岗位"], ["李四", "销售"]])

    page = SplitPage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    # 两个文件的共有列只有「姓名」，但用户可以手打一个只有 a 才有的列
    page.cb_column.addItem("部门")
    page.cb_column.setCurrentText("部门")
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True
    assert [p.name for p in page._header_report.blocked] == ["b.xlsx"]
    assert "缺少「部门」" in page.log.toPlainText()
    page.deleteLater()


def test_对比页的体检只看AB两个文件(qapp, 静音弹窗, 写xlsx, tmp_path):
    """清单里的第三个文件跟这次对比无关，不该被算进体检。"""
    from filebatch.ui.excel_toolbox.compare_page import ComparePage

    src = tmp_path / "源"
    写xlsx(src / "a.xlsx", [["工号", "金额"], ["001", "100"]])
    写xlsx(src / "b.xlsx", [["工号", "金额"], ["001", "200"]])
    写xlsx(src / "无关.xlsx", [["完全不同的列"], ["x"]])

    page = ComparePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page.list_keys.set_checked(["工号"])
    page._output_path = tmp_path / "对比结果.xlsx"

    assert page.generate_preview() is True
    assert page._header_report.total == 2, "只体检 A 和 B"
    assert page._header_report.blocked == []
    page.deleteLater()


def test_不选列时体检只报表头不一致(qapp, 静音弹窗, 写xlsx, tmp_path):
    from filebatch.ui.excel_toolbox.dedupe_page import DedupePage

    src = tmp_path / "源"
    写xlsx(src / "a.xlsx", [["A", "B"], ["1", "2"], ["1", "2"]])
    写xlsx(src / "b.xlsx", [["A", "C"], ["1", "3"], ["1", "3"]])

    page = DedupePage()
    page.file_list.add_paths_sync([src])
    等列名(qapp, page)
    page._output_path = tmp_path / "输出"

    assert page.generate_preview() is True, "整行去重不要求任何列"
    assert page._header_report.missing == []
    assert len(page._header_report.mismatched) == 1
    page.deleteLater()


# ---------------- 格式页的对齐：和字体同一套规矩 ----------------

def test_格式页默认是保留原对齐(qapp, 静音弹窗):
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    page = FormatPage()
    for cb, 名 in ((page.cb_body_align, "正文"), (page.cb_head_align, "表头")):
        assert cb.currentIndex() == 0, f"{名}对齐默认应该是第一项"
        assert "保留原对齐" in cb.currentText(), f"{名}：{cb.currentText()}"
        assert page.selected_align(cb) is None, f"{名}拿到的不是 None"

    opts = page.current_options()
    assert opts.body_align is None and opts.header_align is None
    assert opts.validate() is None, "默认参数必须能过校验"
    page.deleteLater()


def test_格式页默认执行后原对齐不变(qapp, 静音弹窗, tmp_path):
    from openpyxl.styles import Alignment

    from filebatch.ui.excel_toolbox.format_page import FormatPage

    src = tmp_path / "源" / "报表.xlsx"
    src.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r in [["表头"], ["左"], ["中"], ["右"]]:
        ws.append(r)
    for 坐标, 水平 in (("A1", "right"), ("A2", "left"), ("A3", "center"), ("A4", "right")):
        ws[坐标].alignment = Alignment(horizontal=水平)
    wb.save(src)

    page = FormatPage()
    page.file_list.add_paths_sync([src])
    ws = 跑格式页(qapp, page, tmp_path)

    得到 = [ws[f"A{i}"].alignment.horizontal for i in range(1, 5)]
    assert 得到 == ["right", "left", "center", "right"], 得到
    page.deleteLater()


def test_格式页选了对齐才改(qapp, 静音弹窗, tmp_path):
    from openpyxl.styles import Alignment

    from filebatch.ui.excel_toolbox.format_page import FormatPage

    src = tmp_path / "源" / "报表.xlsx"
    src.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["表头"]); ws.append(["正文"])
    ws["A1"].alignment = Alignment(horizontal="left")
    ws["A2"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    wb.save(src)

    page = FormatPage()
    page.file_list.add_paths_sync([src])
    页面项 = [page.cb_body_align.itemText(i) for i in range(page.cb_body_align.count())]
    assert 页面项[1:] == ["左对齐", "居中", "右对齐"], 页面项
    page.cb_body_align.setCurrentIndex(3)      # 右对齐
    assert page.selected_align(page.cb_body_align) == "right"

    ws = 跑格式页(qapp, page, tmp_path)
    assert ws["A2"].alignment.horizontal == "right"
    assert ws["A2"].alignment.vertical == "top", "垂直对齐不该被误伤"
    assert ws["A2"].alignment.wrap_text is True, "自动换行不该被误伤"
    assert ws["A1"].alignment.horizontal == "left", "没设表头对齐就不该动表头"
    page.deleteLater()


def test_格式页表头和正文对齐互不绑死(qapp, 静音弹窗, tmp_path):
    from openpyxl.styles import Alignment

    from filebatch.ui.excel_toolbox.format_page import FormatPage

    src = tmp_path / "源" / "报表.xlsx"
    src.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["表头"]); ws.append(["正文"])
    for 坐标 in ("A1", "A2"):
        ws[坐标].alignment = Alignment(horizontal="left")
    wb.save(src)

    page = FormatPage()
    page.file_list.add_paths_sync([src])
    page.cb_head_align.setCurrentIndex(2)      # 表头居中，正文保持"保留原对齐"
    assert page.selected_align(page.cb_head_align) == "center"
    assert page.selected_align(page.cb_body_align) is None

    ws = 跑格式页(qapp, page, tmp_path)
    assert ws["A1"].alignment.horizontal == "center"
    assert ws["A2"].alignment.horizontal == "left"
    page.deleteLater()


def test_格式页对齐的None不会变成字符串None(qapp, 静音弹窗):
    """Qt 的 currentData() 拿 None 出来，str() 一下就成了非法值 "None"。"""
    from filebatch.ui.excel_toolbox.format_page import FormatPage

    page = FormatPage()
    opts = page.current_options()
    assert opts.body_align != "None", "None 被 str() 成了字符串"
    assert opts.body_align is None
    assert opts.header_align is None
    assert opts.validate() is None
    page.deleteLater()
